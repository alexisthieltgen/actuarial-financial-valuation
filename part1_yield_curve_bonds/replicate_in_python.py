"""Python replication of the Excel workbook (yield curve, bootstrapping, best
estimate, duration/convexity, rate shocks, bond accrued interest and yield).

The original assignment had to be solved with Excel's financial library; this
script independently re-derives every figure of the report as a cross-check.

Usage (from part1_yield_curve_bonds/):  python replicate_in_python.py
"""
from datetime import date
from pathlib import Path

import numpy as np
import openpyxl
from scipy.optimize import brentq, least_squares

HERE = Path(__file__).resolve().parent

# --------------------------------------------------------------------------- #
# Inputs
# --------------------------------------------------------------------------- #
# Banque de France constant-maturity (par) yields on 12/01/2026
MAT = np.array([1, 2, 3, 5, 7, 10, 15, 20, 25, 30], dtype=float)
PAR = np.array([2.13, 2.24, 2.42, 2.79, 3.12, 3.55, 3.95, 4.18, 4.36, 4.46]) / 100

ws = openpyxl.load_workbook(HERE / "data" / "liabilities_cashflows.xlsx", data_only=True).active
premiums = np.array([ws.cell(4, c).value for c in range(2, 32)], dtype=float)
benefits = np.array([ws.cell(5, c).value for c in range(2, 32)], dtype=float)
CF = benefits - premiums  # net liability cash flows, t = 1..30
T = np.arange(1, 31)


# --------------------------------------------------------------------------- #
# Q1 - Nelson-Siegel calibration (RMSE minimisation, c1 > 0)
# --------------------------------------------------------------------------- #
def nelson_siegel(t, b0, b10, b11, c1):
    t = np.asarray(t, dtype=float)
    return (b0 + (b10 / (t * c1)) * (1 - np.exp(-c1 * t))
            + (b11 / (t * c1**2)) * (1 - (c1 * t + 1) * np.exp(-c1 * t)))


# Start from the Excel Solver solution's neighbourhood; bounds enforce c1 > 0
fit = least_squares(lambda x: nelson_siegel(MAT, *x) - PAR,
                    x0=[0.017, 0.002, 0.005, 0.05],
                    bounds=([-1, -1, -1, 1e-6], [1, 1, 1, 5]), xtol=1e-15, ftol=1e-15)
rmse = np.sqrt(np.mean(fit.fun**2))
print(f"[Q1] Re-fit (scipy)  : b0={fit.x[0]:.9f} b10={fit.x[1]:.9f} b11={fit.x[2]:.9f} c1={fit.x[3]:.9f}  RMSE={rmse:.9f}")

# The RMSE surface is very flat around the optimum, so Excel Solver and scipy stop at
# slightly different parameters (RMSE 3.74604e-4 vs 3.74601e-4). To reproduce the report
# to the cent, downstream computations use the Solver parameters of the Excel workbook.
b0, b10, b11, c1 = 0.016841372, 0.001546384, 0.004745281, 0.052446037
rmse_report = np.sqrt(np.mean((nelson_siegel(MAT, b0, b10, b11, c1) - PAR) ** 2))
print(f"[Q1] Report (Solver) : b0={b0:.9f} b10={b10:.9f} b11={b11:.9f} c1={c1:.9f}  RMSE={rmse_report:.9f}")

# Full par curve: observed points where available, NS elsewhere
par_curve = nelson_siegel(T, b0, b10, b11, c1)
for m, y in zip(MAT.astype(int), PAR):
    par_curve[m - 1] = y


# --------------------------------------------------------------------------- #
# Q2 - Bootstrapping of zero-coupon (spot) rates and best estimate
# --------------------------------------------------------------------------- #
spot = np.zeros(30)
for n in range(1, 31):
    c = par_curve[n - 1]
    annuity = np.sum(c / (1 + spot[: n - 1]) ** T[: n - 1])
    spot[n - 1] = ((1 + c) / (1 - annuity)) ** (1 / n) - 1

BE = np.sum(CF / (1 + spot) ** T)
print(f"[Q2] s(1)={spot[0]:.4%}  s(10)={spot[9]:.4%}  s(30)={spot[29]:.4%}")
print(f"[Q2] Best estimate = {BE:,.3f}")


# --------------------------------------------------------------------------- #
# Q3 - Yield of liabilities, modified duration, convexity (no built-ins)
# --------------------------------------------------------------------------- #
yL = brentq(lambda y: np.sum(CF / (1 + y) ** T) - BE, -0.5, 1.0, xtol=1e-14)
pv = CF / (1 + yL) ** T
D = np.sum(T * pv) / BE
D_mod = D / (1 + yL)
convexity = np.sum(T * (T + 1) * pv) / (BE * (1 + yL) ** 2)
print(f"[Q3] yL={yL:.4%}  Macaulay D={D:.6f}  Dmod={D_mod:.6f}  Convexity={convexity:.4f}")


# --------------------------------------------------------------------------- #
# Q4 - Impact of a +/-1% parallel shift
# --------------------------------------------------------------------------- #
print("[Q4] Method                       shift   new BE          rel. change")
for dy in (+0.01, -0.01):
    approx1 = BE * (1 - D_mod * dy)
    approx2 = BE * (1 - D_mod * dy + 0.5 * convexity * dy**2)
    exact = np.sum(CF / (1 + spot + dy) ** T)
    for name, v in (("Modified duration", approx1), ("Duration + convexity", approx2), ("Exact revaluation", exact)):
        print(f"     {name:<28}{dy:+.0%}   {v:>14,.3f}   {v / BE - 1:+.4%}")


# --------------------------------------------------------------------------- #
# Q5/Q6 - OLO 2.75% 22/04/2039: accrued interest, dirty price, yield (ACT/ACT)
# --------------------------------------------------------------------------- #
settle, last_cpn, next_cpn = date(2026, 1, 9), date(2025, 4, 22), date(2026, 4, 22)
coupon, clean, redemption, n_cpn = 2.75, 90.55, 100.0, 14

dcf = (settle - last_cpn).days / (next_cpn - last_cpn).days  # 262 / 365
accrued = coupon * dcf
dirty = clean + accrued
print(f"[Q5] Accrued interest = {accrued:.6f}  Dirty price = {dirty:.6f}")

k = np.arange(1, n_cpn + 1)
price = lambda y: np.sum(coupon / (1 + y) ** (k - dcf)) + redemption / (1 + y) ** (n_cpn - dcf)
yB = brentq(lambda y: price(y) - dirty, 1e-6, 0.5, xtol=1e-14)
print(f"[Q6] Bond yield = {yB:.6%}")
